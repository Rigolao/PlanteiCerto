#!/usr/bin/env python3
"""
Gera o SQL de carga da proveniência por valor (public.tree_value_provenance).

Fontes dos 1.480 registros:
  - aba "Projeções e fontes" da planilha de homologação (447 registros: DAP e altura
    aos 10, 20 e 30 anos, com faixa operacional, origem, confiança, método e fonte);
  - aba "Parametrização 25" (1.025 registros: 41 campos × 25 espécies novas);
  - 8 registros de consolidação do SQL 06 de 03/10/2026 (constantes neste script):
    campos que estavam NULL nos IDs 11, 19 e 79 e foram preenchidos com o valor do
    registro duplicado inativado (82, 97 e 120).

A exportação de produção (CSV de `SELECT * FROM trees`) serve só para validar: cada
valor da planilha tem de ser igual ao que está em public.trees. Qualquer divergência,
formato inesperado, contagem diferente ou chave repetida aborta com exit 1 ANTES de
gravar qualquer arquivo. Não há correção automática de dados.

Uso:
    python3 -E scripts/gerar_sql_proveniencia.py \\
        --planilha <Plantei_Certo_Carga_152_Homologacao_Condicional.xlsx> \\
        --export <CSV do SELECT * FROM trees> \\
        --saida supabase/snippets/provenance-import-20261017/01-importar-proveniencia.sql

`-E` (e não `-I`) porque o openpyxl está instalado no site do usuário.
"""

import argparse
import csv
import math
import os
import re
import sys
from collections import Counter
from decimal import Decimal
from pathlib import Path

ABA_PROJECOES = 'Projeções e fontes'
ABA_PARAMETRIZACAO = 'Parametrização 25'
ABA_CONSOLIDACAO = 'SQL 06 (03/10/2026)'

LINHA_CABECALHO = {ABA_PROJECOES: 4, ABA_PARAMETRIZACAO: 6}

ESPERADO = {ABA_PROJECOES: 447, ABA_PARAMETRIZACAO: 1025, ABA_CONSOLIDACAO: 8}

ORIGENS = (
    'published_table',
    'published_equation',
    'published_description',
    'estimated',
    'estimated_rubric',
    'estimated_taxon_proxy',
    'retido_da_fase_anterior',
    'consolidado_de_registro_inativo',
    'sem_dado_admissivel',
)

# Confianças da "Parametrização 25" gravadas como NULL: o significado já está na origem
# (sem_dado_admissivel → não_avaliada; retido_da_fase_anterior → conforme_fase_anterior).
CONFIANCAS_NULAS = {'não_avaliada', 'conforme_fase_anterior'}

STATUS_ID_PROPOSTO = 'ID PROPOSTO — não importar sem cadastro'

TOLERANCIA = Decimal('1e-6')

# Colunas de public.trees em produção (63) e como comparar cada uma.
NUM, BOOL, TEXTO = 'numero', 'booleano', 'texto'
TIPOS_TREES = {
    'id': NUM,
    'foto': TEXTO,
    'nome_cientifico': TEXTO,
    'nome_popular': TEXTO,
    'origem': TEXTO,
    'decidua_perenifolia': TEXTO,
    'epoca_floracao': TEXTO,
    'epoca_frutificacao': TEXTO,
    'altura_adulta_max_m': NUM,
    'porte_altura_classe': TEXTO,
    'diametro_copa_adulto_max_m': NUM,
    'copa_classe': TEXTO,
    'dap_adulto_max_cm': NUM,
    'altura_primeira_bifurcacao_m': TEXTO,
    'forma_copa': TEXTO,
    'faixa_serv_min_m_recomendada': NUM,
    'berco_area_min_m2_recomendada': NUM,
    'volume_solo_min_m3_recomendado': NUM,
    'compat_fiacao': TEXTO,
    'potencial_dano_calcada_1a5': NUM,
    'tolerancia_sol_pleno': BOOL,
    'tolerancia_meia_sombra': BOOL,
    'tolerancia_sombra': BOOL,
    'tolerancia_seca_1a5': NUM,
    'tolerancia_encharcamento_1a5': NUM,
    'tolerancia_poluicao_atmosferica_1a5': NUM,
    'tolerancia_compactacao_solo_1a5': NUM,
    'tolerancia_ventos_fortes_1a5': NUM,
    'potencial_sujeira_1a5': NUM,
    'presenca_espinhos': BOOL,
    'presenca_subst_irritantes': BOOL,
    'atracao_fauna_1a5': NUM,
    'tolerancia_poda_1a5': NUM,
    'potencial_sombra_1a5': NUM,
    'contribuicao_biodiversidade_1a5': NUM,
    'ativa': BOOL,
    'dap_10a_cm': NUM,
    'dap_10a_min_cm': NUM,
    'dap_10a_max_cm': NUM,
    'altura_10a_m': NUM,
    'altura_10a_min_m': NUM,
    'altura_10a_max_m': NUM,
    'biomassa_aerea_10a_kg': NUM,
    'carbono_armazenado_10a_kg': NUM,
    'co2e_10a_kg': NUM,
    'sobrevivencia_10a_pct': NUM,
    'co2e_esperado_por_muda_10a_kg': NUM,
    'classe_bvoc': TEXTO,
    'evidencia_bvoc': TEXTO,
    'confianca_bvoc': TEXTO,
    'exibir_aviso_bvoc': BOOL,
    'dap_20a_cm': NUM,
    'dap_20a_min_cm': NUM,
    'dap_20a_max_cm': NUM,
    'altura_20a_m': NUM,
    'altura_20a_min_m': NUM,
    'altura_20a_max_m': NUM,
    'dap_30a_cm': NUM,
    'dap_30a_min_cm': NUM,
    'dap_30a_max_cm': NUM,
    'altura_30a_m': NUM,
    'altura_30a_min_m': NUM,
    'altura_30a_max_m': NUM,
}

UNIDADE_POR_VARIAVEL = {'dap': 'cm', 'altura': 'm'}
IDADES = (10, 20, 30)

# Consolidação do SQL 06 de 03/10/2026: (destino, origem inativada, nome científico da
# origem, campo). Os nomes vêm do SQL 04 de 03/10/2026 e são conferidos na exportação.
CONSOLIDACAO = (
    (11, 82, 'Platypodium elegans', 'atracao_fauna_1a5'),
    (11, 82, 'Platypodium elegans', 'potencial_sombra_1a5'),
    (11, 82, 'Platypodium elegans', 'contribuicao_biodiversidade_1a5'),
    (19, 97, 'Calyptranthes clusiifolia', 'epoca_floracao'),
    (19, 97, 'Calyptranthes clusiifolia', 'epoca_frutificacao'),
    (19, 97, 'Calyptranthes clusiifolia', 'altura_primeira_bifurcacao_m'),
    (79, 120, 'Senna fistula (L.) H.S.Irwin & Barneby', 'berco_area_min_m2_recomendada'),
    (79, 120, 'Senna fistula (L.) H.S.Irwin & Barneby', 'volume_solo_min_m3_recomendado'),
)
METODO_CONSOLIDACAO = (
    'Campo NULL no destino preenchido com o valor do registro duplicado inativado; '
    'valores já preenchidos no destino foram preservados'
)

COLUNAS_INSERT = (
    'tree_id', 'campo', 'variavel', 'idade_anos', 'valor_registrado', 'unidade',
    'min_operacional', 'max_operacional', 'origem', 'confianca', 'metodo', 'fonte',
    'status_dominio', 'observacao', 'aba_origem', 'chave_canonica',
)
COLUNAS_INTEIRAS = {'tree_id', 'idade_anos'}
COLUNAS_NUMERICAS = {'min_operacional', 'max_operacional'}

MSG_REEXECUCAO = 'Carga de proveniência já aplicada'


class FormatoInesperado(Exception):
    pass


# ---------------------------------------------------------------------------
# Normalização
# ---------------------------------------------------------------------------

NUMERO_PLANILHA = re.compile(r'^[+-]?\d+(?:[.,]\d+)?$')
NUMERO_EXPORT = re.compile(r'^-?\d+(?:\.\d+)?$')


def para_decimal(v):
    """Valor numérico da planilha (int, float ou texto com vírgula/ponto) → Decimal."""
    if v is None:
        return None
    if isinstance(v, bool):
        raise FormatoInesperado(f'booleano {v!r} onde se esperava número')
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        if not math.isfinite(v):
            raise FormatoInesperado(f'número não finito {v!r}')
        return Decimal(repr(v))
    if isinstance(v, str) and NUMERO_PLANILHA.match(v.strip()):
        return Decimal(v.strip().replace(',', '.'))
    raise FormatoInesperado(f'{v!r} não é número')


def para_bool(v):
    """Booleano da planilha: True/False, 1/0 ou texto sim/não/true/false/1/0."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int) and v in (0, 1):
        return bool(v)
    if isinstance(v, str):
        s = v.strip().lower()
        if s in ('sim', 'true', '1'):
            return True
        if s in ('não', 'false', '0'):
            return False
    raise FormatoInesperado(f'{v!r} não é booleano')


def para_texto(v):
    """Texto da planilha com espaços normalizados (strip + espaços internos únicos)."""
    if v is None:
        return None
    if isinstance(v, str):
        return ' '.join(v.split())
    raise FormatoInesperado(f'{v!r} ({type(v).__name__}) onde se esperava texto')


def normalizar_planilha(tipo, v):
    if tipo == NUM:
        return para_decimal(v)
    if tipo == BOOL:
        return para_bool(v)
    return para_texto(v)


def normalizar_export(tipo, s):
    """Valor do CSV exportado (NULL escrito como `null`)."""
    if s == 'null':
        return None
    if tipo == NUM:
        if not NUMERO_EXPORT.match(s):
            raise FormatoInesperado(f'{s!r} não é número')
        return Decimal(s)
    if tipo == BOOL:
        if s == 'true':
            return True
        if s == 'false':
            return False
        raise FormatoInesperado(f'{s!r} não é booleano')
    return ' '.join(s.split())


def iguais(tipo, a, b):
    if a is None or b is None:
        return a is None and b is None
    if tipo == NUM:
        return abs(a - b) <= TOLERANCIA
    return a == b


def texto_decimal(d):
    return None if d is None else format(d, 'f')


def para_inteiro(v):
    if isinstance(v, bool):
        raise FormatoInesperado(f'booleano {v!r} onde se esperava inteiro')
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and re.fullmatch(r'\s*\d+\s*', v):
        return int(v)
    raise FormatoInesperado(f'{v!r} não é inteiro')


def texto_livre(v):
    """Metadado textual (método, fonte, observação...): gravado como está na planilha."""
    if v is None:
        return None
    if isinstance(v, str):
        return v
    raise FormatoInesperado(f'{v!r} ({type(v).__name__}) onde se esperava texto')


def exibir(v):
    if v is None:
        return 'NULL'
    if isinstance(v, Decimal):
        return format(v, 'f')
    return repr(v)


# ---------------------------------------------------------------------------
# Leitura
# ---------------------------------------------------------------------------

def ler_export(caminho, problemas):
    with open(caminho, newline='', encoding='utf-8-sig') as f:
        leitor = csv.reader(f)
        cabecalho = next(leitor, None)
        if cabecalho is None:
            problemas.append(f'exportação {caminho}: arquivo vazio')
            return {}
        repetidas = [c for c, n in Counter(cabecalho).items() if n > 1]
        if repetidas:
            problemas.append(f'exportação: colunas repetidas {repetidas}')
        faltando = [c for c in TIPOS_TREES if c not in cabecalho]
        sobrando = [c for c in cabecalho if c not in TIPOS_TREES]
        if faltando or sobrando:
            problemas.append(
                f'exportação: colunas diferentes do esperado (faltando {faltando}, a mais {sobrando})'
            )
            return {}
        arvores = {}
        for n_linha, linha in enumerate(leitor, start=2):
            if len(linha) != len(cabecalho):
                problemas.append(
                    f'exportação linha {n_linha}: {len(linha)} campos; esperado {len(cabecalho)}'
                )
                continue
            registro = dict(zip(cabecalho, linha))
            try:
                tid = int(registro['id'])
            except ValueError:
                problemas.append(f'exportação linha {n_linha}: id {registro["id"]!r} não é inteiro')
                continue
            if tid in arvores:
                problemas.append(f'exportação: ID {tid} repetido')
                continue
            arvores[tid] = registro
    return arvores


def ler_aba(livro, aba, colunas, problemas):
    if aba not in livro.sheetnames:
        problemas.append(f'planilha: aba "{aba}" não encontrada')
        return []
    linhas = list(livro[aba].iter_rows(values_only=True))
    n_cab = LINHA_CABECALHO[aba]
    if len(linhas) < n_cab:
        problemas.append(f'aba "{aba}": sem a linha de cabeçalho {n_cab}')
        return []
    cabecalho = [c.strip() if isinstance(c, str) else c for c in linhas[n_cab - 1]]
    indice = {}
    for nome in colunas:
        if cabecalho.count(nome) != 1:
            problemas.append(
                f'aba "{aba}": coluna "{nome}" não encontrada (ou repetida) na linha {n_cab}'
            )
        else:
            indice[nome] = cabecalho.index(nome)
    if len(indice) != len(colunas):
        return []
    registros = []
    for n_linha, linha in enumerate(linhas[n_cab:], start=n_cab + 1):
        if all(v is None for v in linha):
            continue
        valores = {nome: (linha[i] if i < len(linha) else None) for nome, i in indice.items()}
        registros.append((n_linha, valores))
    return registros


COLUNAS_PROJECOES = (
    'id_proposto', 'canonical_key', 'nome_cientifico', 'variavel', 'idade_anos', 'valor',
    'min_operacional', 'max_operacional', 'unidade', 'origem', 'confianca', 'metodo',
    'fonte', 'status_dominio', 'observacao', 'status_carga',
)
COLUNAS_PARAMETRIZACAO = (
    'ID', 'Espécie', 'Campo do app', 'Valor proposto', 'Origem', 'Confiança',
    'Método / base de inferência', 'Fonte', 'Observação / ressalva',
)


# ---------------------------------------------------------------------------
# Validação e montagem dos registros
# ---------------------------------------------------------------------------

class Validador:
    def __init__(self, arvores):
        self.arvores = arvores
        self.problemas = []
        self.ids_conferidos = set()
        self.nomes_conferidos = set()

    def erro(self, msg):
        self.problemas.append(msg)

    def conferir_id(self, tid, onde):
        if tid in self.ids_conferidos:
            return tid in self.arvores and self.arvores[tid]['ativa'] == 'true'
        self.ids_conferidos.add(tid)
        if tid not in self.arvores:
            self.erro(f'{tid}: ID não existe na exportação ({onde})')
            return False
        if self.arvores[tid]['ativa'] != 'true':
            self.erro(f'{tid}: ID não está ativo na exportação ({onde})')
            return False
        return True

    def conferir_nome(self, tid, nome, aba):
        if (tid, aba) in self.nomes_conferidos or tid not in self.arvores:
            return
        self.nomes_conferidos.add((tid, aba))
        try:
            planilha = para_texto(nome)
        except FormatoInesperado as e:
            self.erro(f'{tid}: nome científico em "{aba}" com formato inesperado: {e}')
            return
        export = normalizar_export(TEXTO, self.arvores[tid]['nome_cientifico'])
        if planilha != export:
            self.erro(
                f'{tid}: nome científico em "{aba}" {planilha!r} difere da exportação {export!r}'
            )

    def comparar(self, tid, campo, valor_planilha, tipo=None):
        """Compara um valor já normalizado com a coluna `campo` da exportação."""
        tipo = tipo or TIPOS_TREES[campo]
        try:
            export = normalizar_export(tipo, self.arvores[tid][campo])
        except FormatoInesperado as e:
            self.erro(f'{tid} {campo}: formato inesperado na exportação: {e}')
            return
        if not iguais(tipo, valor_planilha, export):
            self.erro(
                f'divergência {tid} {campo}: planilha={exibir(valor_planilha)} '
                f'exportação={exibir(export)}'
            )

    def conferir_origem(self, tid, campo, origem):
        if origem not in ORIGENS:
            self.erro(f'{tid} {campo}: origem {origem!r} fora das 9 chaves da tabela')


def montar_projecoes(registros, v):
    saida = []
    for n_linha, r in registros:
        onde = f'"{ABA_PROJECOES}" linha {n_linha}'
        try:
            tid = para_inteiro(r['id_proposto'])
            idade = para_inteiro(r['idade_anos'])
            variavel = para_texto(r['variavel'])
            unidade = para_texto(r['unidade'])
            valor = para_decimal(r['valor'])
            minimo = para_decimal(r['min_operacional'])
            maximo = para_decimal(r['max_operacional'])
            origem = para_texto(r['origem'])
            metadados = {
                'chave_canonica': texto_livre(r['canonical_key']),
                'confianca': texto_livre(r['confianca']),
                'metodo': texto_livre(r['metodo']),
                'fonte': texto_livre(r['fonte']),
                'status_dominio': texto_livre(r['status_dominio']),
                'observacao': texto_livre(r['observacao']),
            }
            status_carga = para_texto(r['status_carga'])
        except FormatoInesperado as e:
            v.erro(f'{onde}: formato inesperado: {e}')
            continue
        if variavel not in UNIDADE_POR_VARIAVEL or idade not in IDADES:
            v.erro(f'{onde}: variável/idade inesperadas ({variavel!r}, {idade!r})')
            continue
        sufixo = UNIDADE_POR_VARIAVEL[variavel]
        campo = f'{variavel}_{idade}a_{sufixo}'
        col_min = f'{variavel}_{idade}a_min_{sufixo}'
        col_max = f'{variavel}_{idade}a_max_{sufixo}'
        if unidade != sufixo:
            v.erro(f'{tid} {campo}: unidade {unidade!r}; esperado {sufixo!r}')
        v.conferir_origem(tid, campo, origem)
        if v.conferir_id(tid, onde):
            v.conferir_nome(tid, r['nome_cientifico'], ABA_PROJECOES)
            v.comparar(tid, campo, valor)
            v.comparar(tid, col_min, minimo)
            v.comparar(tid, col_max, maximo)
        saida.append({
            'tree_id': tid,
            'campo': campo,
            'variavel': variavel,
            'idade_anos': idade,
            'valor_registrado': texto_decimal(valor),
            'unidade': unidade,
            'min_operacional': minimo,
            'max_operacional': maximo,
            'origem': origem,
            'aba_origem': ABA_PROJECOES,
            # usados só na trava e no cabeçalho
            '_col_min': col_min,
            '_col_max': col_max,
            '_valor': valor,
            '_status_carga': status_carga,
            **metadados,
        })
    return saida


def valor_registrado_parametrizacao(tipo, normalizado, bruto):
    if normalizado is None:
        return None
    if tipo == NUM:
        return texto_decimal(normalizado)
    if tipo == BOOL:
        return 'true' if normalizado else 'false'
    return bruto


def montar_parametrizacao(registros, v):
    saida = []
    for n_linha, r in registros:
        onde = f'"{ABA_PARAMETRIZACAO}" linha {n_linha}'
        try:
            tid = para_inteiro(r['ID'])
            campo = para_texto(r['Campo do app'])
            origem = para_texto(r['Origem'])
            confianca = texto_livre(r['Confiança'])
            metadados = {
                'metodo': texto_livre(r['Método / base de inferência']),
                'fonte': texto_livre(r['Fonte']),
                'observacao': texto_livre(r['Observação / ressalva']),
            }
        except FormatoInesperado as e:
            v.erro(f'{onde}: formato inesperado: {e}')
            continue
        tipo = TIPOS_TREES.get(campo)
        if tipo is None or campo in ('id', 'ativa'):
            v.erro(f'{tid} {campo}: campo não é atributo de public.trees ({onde})')
            continue
        bruto = r['Valor proposto']
        try:
            valor = normalizar_planilha(tipo, bruto)
        except FormatoInesperado as e:
            v.erro(f'{tid} {campo}: formato inesperado na planilha: {e}')
            continue
        v.conferir_origem(tid, campo, origem)
        if v.conferir_id(tid, onde):
            v.conferir_nome(tid, r['Espécie'], ABA_PARAMETRIZACAO)
            if origem == 'sem_dado_admissivel':
                if valor is not None:
                    v.erro(f'{tid} {campo}: sem_dado_admissivel com valor na planilha ({bruto!r})')
                export = v.arvores[tid][campo]
                if export != 'null':
                    v.erro(f'{tid} {campo}: sem_dado_admissivel exige NULL na exportação; há {export!r}')
            else:
                v.comparar(tid, campo, valor, tipo)
        saida.append({
            'tree_id': tid,
            'campo': campo,
            'variavel': None,
            'idade_anos': None,
            'valor_registrado': valor_registrado_parametrizacao(tipo, valor, bruto),
            'unidade': None,
            'min_operacional': None,
            'max_operacional': None,
            'origem': origem,
            'confianca': None if confianca in CONFIANCAS_NULAS else confianca,
            'status_dominio': None,
            'aba_origem': ABA_PARAMETRIZACAO,
            'chave_canonica': None,
            **metadados,
        })
    return saida


def montar_consolidacao(v):
    saida = []
    for destino, origem_id, nome_origem, campo in CONSOLIDACAO:
        onde = f'consolidação {origem_id}→{destino}'
        if not v.conferir_id(destino, onde):
            continue
        if origem_id not in v.arvores:
            v.erro(f'{origem_id}: registro de origem da consolidação não existe na exportação')
            continue
        fonte_reg = v.arvores[origem_id]
        if fonte_reg['nome_cientifico'] != nome_origem:
            v.erro(
                f'{origem_id}: nome científico {fonte_reg["nome_cientifico"]!r} difere do '
                f'esperado {nome_origem!r} ({onde})'
            )
        if fonte_reg['ativa'] != 'false':
            v.erro(f'{origem_id}: registro de origem da consolidação não está inativo')
        tipo = TIPOS_TREES[campo]
        valor_destino = v.arvores[destino][campo]
        if valor_destino == 'null':
            v.erro(f'{destino} {campo}: valor consolidado está NULL na exportação')
        else:
            try:
                iguais_origem = iguais(
                    tipo,
                    normalizar_export(tipo, valor_destino),
                    normalizar_export(tipo, fonte_reg[campo]),
                )
            except FormatoInesperado as e:
                v.erro(f'{destino} {campo}: formato inesperado na exportação: {e}')
                iguais_origem = True
            if not iguais_origem:
                v.erro(
                    f'divergência {destino} {campo}: destino={valor_destino!r} difere da '
                    f'origem {origem_id}={fonte_reg[campo]!r}'
                )
        saida.append({
            'tree_id': destino,
            'campo': campo,
            'variavel': None,
            'idade_anos': None,
            'valor_registrado': None if valor_destino == 'null' else valor_destino,
            'unidade': None,
            'min_operacional': None,
            'max_operacional': None,
            'origem': 'consolidado_de_registro_inativo',
            'confianca': None,
            'metodo': METODO_CONSOLIDACAO,
            'fonte': f'Registro inativo ID {origem_id} ({nome_origem})',
            'status_dominio': None,
            'observacao': None,
            'aba_origem': ABA_CONSOLIDACAO,
            'chave_canonica': None,
        })
    return saida


def esperado_no_banco(r, arvores):
    """(tipo, valor) que a trava 2b espera encontrar em public.trees para o registro.

    - sem_dado_admissivel ou NULL na exportação: 'nulo' (o campo tem de estar NULL);
    - coluna numérica: 'numero', comparado como numeric com o valor registrado;
    - booleano e texto: o texto bruto do banco na exportação conferida, comparado por
      igualdade exata com `to_jsonb(t) ->> campo`.
    """
    bruto = arvores[r['tree_id']][r['campo']]
    if r['origem'] == 'sem_dado_admissivel' or bruto == 'null':
        return 'nulo', None
    tipo = TIPOS_TREES[r['campo']]
    if tipo == NUM:
        return 'numero', r['valor_registrado']
    if tipo == BOOL:
        return 'booleano', bruto
    return 'texto', bruto


def conferir_conjunto(fontes, v):
    for aba, registros in fontes:
        if len(registros) != ESPERADO[aba]:
            v.erro(f'aba "{aba}": {len(registros)} registros; esperado {ESPERADO[aba]}')
    vistos = {}
    for aba, registros in fontes:
        for r in registros:
            chave = (r['tree_id'], r['campo'])
            if chave in vistos:
                v.erro(f'{chave[0]} {chave[1]}: chave repetida ("{vistos[chave]}" e "{aba}")')
            else:
                vistos[chave] = aba


# ---------------------------------------------------------------------------
# SQL
# ---------------------------------------------------------------------------

def sql_texto(s):
    if s is None:
        return 'NULL'
    if '\x00' in s:
        raise FormatoInesperado(f'caractere NUL em {s!r}')
    return "'" + s.replace("'", "''") + "'"


def sql_numero(d):
    return 'NULL' if d is None else format(d, 'f')


def sql_valor(coluna, v):
    if coluna in COLUNAS_INTEIRAS:
        return 'NULL' if v is None else str(int(v))
    if coluna in COLUNAS_NUMERICAS:
        return sql_numero(v)
    return sql_texto(v)


def milhar(n):
    return f'{n:,}'.replace(',', '.')


def lista_sql(itens):
    return ', '.join(sql_texto(i) for i in itens)


def bloco_insert(titulo, registros):
    linhas = [
        f'-- {titulo}: {milhar(len(registros))} registros.',
        'INSERT INTO public.tree_value_provenance (',
        '  ' + ', '.join(COLUNAS_INSERT),
        ') VALUES',
    ]
    tuplas = [
        '  (' + ', '.join(sql_valor(c, r[c]) for c in COLUNAS_INSERT) + ')'
        for r in registros
    ]
    linhas.append(',\n'.join(tuplas) + ';')
    return '\n'.join(linhas)


def gerar_sql(projecoes, parametrizacao, consolidacao, arvores, planilha, export):
    todos = projecoes + parametrizacao + consolidacao
    atributos = parametrizacao + consolidacao
    por_origem = Counter(r['origem'] for r in todos)
    ids_carga = sorted({r['tree_id'] for r in todos})
    abas = (ABA_PROJECOES, ABA_PARAMETRIZACAO, ABA_CONSOLIDACAO)
    proposto = [r for r in projecoes if r['_status_carga'] == STATUS_ID_PROPOSTO]
    ids_proposto = sorted({r['tree_id'] for r in proposto})
    faixa_proposto = (
        f'{ids_proposto[0]}–{ids_proposto[-1]}' if ids_proposto else 'nenhum'
    )
    colunas_projecao = sorted(
        {r['campo'] for r in projecoes}
        | {r['_col_min'] for r in projecoes}
        | {r['_col_max'] for r in projecoes}
    )
    por_variavel_idade = Counter((r['variavel'], r['idade_anos']) for r in projecoes)
    ids_param = sorted({r['tree_id'] for r in parametrizacao})
    faixa_param = f'{ids_param[0]}–{ids_param[-1]}' if ids_param else 'nenhum'
    n_campos_param = len({r['campo'] for r in parametrizacao})

    origens_txt = '\n'.join(
        f'--   {o:<32} {milhar(por_origem[o]):>5}' for o in ORIGENS if por_origem[o]
    )
    projecoes_txt = ', '.join(
        f'{var} {idade}a {por_variavel_idade[(var, idade)]}'
        for var in ('dap', 'altura') for idade in IDADES
    )

    cabecalho = f"""\
-- Carga da proveniência por valor do catálogo de árvores — 17/10/2026.
--
-- GERADO por scripts/gerar_sql_proveniencia.py; não edite à mão. Para regenerar:
--   python3 -E scripts/gerar_sql_proveniencia.py --planilha <xlsx> --export <csv> \\
--     --saida supabase/snippets/provenance-import-20261017/01-importar-proveniencia.sql
--
-- Fontes:
--   Planilha {Path(planilha).name}
--     aba "{ABA_PROJECOES}" (cabeçalho na linha 4): {milhar(len(projecoes))} registros
--       ({projecoes_txt})
--     aba "{ABA_PARAMETRIZACAO}" (cabeçalho na linha 6): {milhar(len(parametrizacao))} registros
--       ({n_campos_param} campos × {len(ids_param)} espécies novas, IDs {faixa_param})
--   SQL 06 de 03/10/2026 (trees-catalog-import-20261003/06-complementar-fichas-arvores.sql):
--     {len(consolidacao)} campos copiados de registros inativados (82→11, 97→19, 120→79)
--   Conferido pelo gerador contra a exportação de produção {Path(export).name}
--   (SELECT * FROM trees): 0 divergências de valor, mínimo e máximo; as
--   {milhar(por_origem['sem_dado_admissivel'])} linhas sem_dado_admissivel estão NULL em public.trees.
--
-- Total: {milhar(len(todos))} linhas em public.tree_value_provenance. Por origem:
{origens_txt}
--
-- Ressalva da planilha: a coluna status_carga de "{ABA_PROJECOES}" traz
-- "{STATUS_ID_PROPOSTO}" em {len(proposto)} linhas (IDs {faixa_proposto}).
-- Esses {len(ids_proposto)} IDs já estão cadastrados e ativos em public.trees (conferido na
-- exportação de produção), então a ressalva não bloqueia esta carga. status_carga não
-- é importada.
--
-- Confiança: na aba "{ABA_PARAMETRIZACAO}", 'não_avaliada' e 'conforme_fase_anterior'
-- são gravadas como NULL; o significado já está na origem. min_operacional e
-- max_operacional são bandas operacionais de incerteza, não intervalos de confiança.
--
-- Ordem de execução em produção (uma sessão por arquivo; recomendado
-- `psql -v ON_ERROR_STOP=1 -f`, com o SQL Editor como alternativa; ver o README):
--   1. supabase/migrations/20261017000100_create_tree_value_provenance.sql
--   2. este arquivo (01-importar-proveniencia.sql)
--   3. 02-posflight.sql
--
-- Travas (cada uma aborta a transação; nada é gravado):
--   - reexecução: a tabela já tem linhas das abas desta carga;
--   - projeções: valor, mínimo ou máximo de alguma das {milhar(len(projecoes))} projeções difere de
--     public.trees (comparação numérica exata, IS DISTINCT FROM);
--   - parametrização e consolidação: algum dos {milhar(len(atributos))} valores difere de public.trees
--     como estava na exportação conferida (número: comparação numérica exata; booleano e
--     texto: igualdade exata do texto; sem_dado_admissivel: o campo tem de estar NULL);
--   - catálogo: public.trees diferente de 192 árvores e 152 ativas, ou algum dos
--     {len(ids_carga)} IDs desta carga inativo;
--   - conferência final das contagens inseridas por aba.
-- Não altera public.trees.
--
-- Reversão:
--   DELETE FROM public.tree_value_provenance
--   WHERE aba_origem IN ({lista_sql(abas)});
"""

    trava_reexecucao = f"""\
-- Trava 1: reexecução.
DO $$
DECLARE
  existentes integer;
BEGIN
  IF to_regclass('public.tree_value_provenance') IS NULL THEN
    RAISE EXCEPTION 'Carga de proveniência interrompida: public.tree_value_provenance não existe. Aplique antes a migration 20261017000100_create_tree_value_provenance.sql.';
  END IF;

  SELECT count(*) INTO existentes
  FROM public.tree_value_provenance
  WHERE aba_origem IN ({lista_sql(abas)});

  IF existentes > 0 THEN
    RAISE EXCEPTION '{MSG_REEXECUCAO}: public.tree_value_provenance já contém % linha(s) das abas desta carga. Nada foi gravado.', existentes;
  END IF;
END $$;
"""

    tuplas_proj = ',\n'.join(
        f'      ({r["tree_id"]}, {sql_texto(r["campo"])}, {sql_texto(r["_col_min"])}, '
        f'{sql_texto(r["_col_max"])}, {sql_numero(r["_valor"])}, '
        f'{sql_numero(r["min_operacional"])}, {sql_numero(r["max_operacional"])})'
        for r in projecoes
    )
    trava_projecoes = f"""\
-- Trava 2: as {milhar(len(projecoes))} projeções (valor, mínimo e máximo) têm de ser iguais a public.trees.
DO $$
DECLARE
  ausentes integer;
  divergencias integer;
BEGIN
  SELECT count(*) INTO ausentes
  FROM unnest(ARRAY[{lista_sql(colunas_projecao)}]::text[]) AS esperado(coluna)
  LEFT JOIN information_schema.columns c
    ON c.table_schema = 'public'
   AND c.table_name = 'trees'
   AND c.column_name = esperado.coluna
  WHERE c.column_name IS NULL OR c.data_type <> 'numeric';

  IF ausentes <> 0 THEN
    RAISE EXCEPTION 'Carga de proveniência interrompida: % coluna(s) de projeção ausente(s) ou não numeric em public.trees. Nada foi gravado.', ausentes;
  END IF;

  SELECT count(*) INTO divergencias
  FROM (VALUES
{tuplas_proj}
  ) AS p(tree_id, coluna_valor, coluna_min, coluna_max, valor, minimo, maximo)
  LEFT JOIN public.trees t ON t.id = p.tree_id
  WHERE t.id IS NULL
     OR (to_jsonb(t) ->> p.coluna_valor)::numeric IS DISTINCT FROM p.valor::numeric
     OR (to_jsonb(t) ->> p.coluna_min)::numeric IS DISTINCT FROM p.minimo::numeric
     OR (to_jsonb(t) ->> p.coluna_max)::numeric IS DISTINCT FROM p.maximo::numeric;

  IF divergencias <> 0 THEN
    RAISE EXCEPTION 'Carga de proveniência interrompida: % de {len(projecoes)} projeções divergem de public.trees (valor, mínimo ou máximo). Nada foi gravado.', divergencias;
  END IF;
END $$;
"""

    tuplas_atributos = []
    for r in atributos:
        tipo_trava, valor_trava = esperado_no_banco(r, arvores)
        tuplas_atributos.append(
            f'      ({r["tree_id"]}, {sql_texto(r["campo"])}, {sql_texto(tipo_trava)}, '
            f'{sql_texto(valor_trava)})'
        )
    tuplas_atributos_txt = ',\n'.join(tuplas_atributos)
    trava_atributos = f"""\
-- Trava 2b: os {milhar(len(atributos))} valores da parametrização e da consolidação têm de estar em
-- public.trees como na exportação conferida pelo gerador. Sem esta trava, um valor
-- editado entre a exportação e a carga ficaria atrelado à fonte antiga sem
-- valor_alterado_em.
DO $$
DECLARE
  divergencias integer;
  exemplos text;
BEGIN
  WITH esperado(tree_id, campo, tipo, valor) AS (
    VALUES
{tuplas_atributos_txt}
  ),
  divergentes AS (
    SELECT e.tree_id, e.campo
    FROM esperado e
    LEFT JOIN public.trees t ON t.id = e.tree_id
    WHERE t.id IS NULL
       OR CASE e.tipo
            WHEN 'nulo' THEN (to_jsonb(t) -> e.campo) IS DISTINCT FROM 'null'::jsonb
            WHEN 'numero' THEN (to_jsonb(t) ->> e.campo)::numeric IS DISTINCT FROM e.valor::numeric
            ELSE (to_jsonb(t) ->> e.campo) IS DISTINCT FROM e.valor
          END
  )
  SELECT (SELECT count(*) FROM divergentes),
         (SELECT string_agg(d.tree_id || ' ' || d.campo, ', ')
            FROM (SELECT tree_id, campo FROM divergentes ORDER BY tree_id, campo LIMIT 10) AS d)
    INTO divergencias, exemplos;

  IF divergencias <> 0 THEN
    RAISE EXCEPTION 'Carga de proveniência interrompida: % de {len(atributos)} valores da parametrização e da consolidação divergem de public.trees (exemplos: %). Nada foi gravado.', divergencias, exemplos;
  END IF;
END $$;
"""

    ids_txt = ', '.join(str(i) for i in ids_carga)
    trava_catalogo = f"""\
-- Trava 3: catálogo com 192 árvores e 152 ativas; os {len(ids_carga)} IDs desta carga ativos.
DO $$
DECLARE
  total integer;
  ativas integer;
  ids_ativos integer;
BEGIN
  SELECT count(*), count(*) FILTER (WHERE ativa) INTO total, ativas
  FROM public.trees;

  IF total <> 192 OR ativas <> 152 THEN
    RAISE EXCEPTION 'Carga de proveniência interrompida: public.trees tem % árvores e % ativas; esperado 192 e 152. Nada foi gravado.', total, ativas;
  END IF;

  SELECT count(*) INTO ids_ativos
  FROM public.trees
  WHERE ativa IS TRUE
    AND id IN ({ids_txt});

  IF ids_ativos <> {len(ids_carga)} THEN
    RAISE EXCEPTION 'Carga de proveniência interrompida: só % dos {len(ids_carga)} IDs desta carga estão ativos. Nada foi gravado.', ids_ativos;
  END IF;
END $$;
"""

    esperado_aba = ', '.join(
        f'({sql_texto(aba)}, {n})'
        for aba, n in (
            (ABA_PROJECOES, len(projecoes)),
            (ABA_PARAMETRIZACAO, len(parametrizacao)),
            (ABA_CONSOLIDACAO, len(consolidacao)),
        )
    )
    conferencia = f"""\
-- Conferência final antes do COMMIT: contagem inserida por aba.
DO $$
DECLARE
  divergentes integer;
BEGIN
  SELECT count(*) INTO divergentes
  FROM (VALUES {esperado_aba}) AS esperado(aba_origem, quantidade)
  WHERE esperado.quantidade <> (
    SELECT count(*) FROM public.tree_value_provenance p
    WHERE p.aba_origem = esperado.aba_origem
  );

  IF divergentes <> 0 THEN
    RAISE EXCEPTION 'Carga de proveniência interrompida: contagem inserida difere do esperado em % aba(s). Nada foi gravado.', divergentes;
  END IF;

  RAISE NOTICE 'Carga de proveniência: {len(todos)} registros inseridos ({len(projecoes)} projeções, {len(parametrizacao)} parametrização, {len(consolidacao)} consolidação).';
END $$;
"""

    partes = [
        cabecalho,
        'BEGIN;',
        '',
        trava_reexecucao,
        trava_projecoes,
        trava_atributos,
        trava_catalogo,
        bloco_insert(f'Aba "{ABA_PROJECOES}"', projecoes) + '\n',
        bloco_insert(f'Aba "{ABA_PARAMETRIZACAO}"', parametrizacao) + '\n',
        bloco_insert(f'Consolidação ({ABA_CONSOLIDACAO})', consolidacao) + '\n',
        conferencia,
        'COMMIT;',
        '',
    ]
    return '\n'.join(partes)


# ---------------------------------------------------------------------------

def main(argv=None):
    parser = argparse.ArgumentParser(
        description='Gera o SQL de carga de public.tree_value_provenance.'
    )
    parser.add_argument('--planilha', required=True, help='xlsx de homologação')
    parser.add_argument('--export', required=True, help='CSV de SELECT * FROM trees (produção)')
    parser.add_argument('--saida', required=True, help='arquivo .sql a gerar')
    args = parser.parse_args(argv)

    for caminho in (args.planilha, args.export):
        if not os.path.isfile(caminho):
            print(f'ERRO: arquivo não encontrado: {caminho}', file=sys.stderr)
            return 1

    try:
        import openpyxl
    except ImportError:
        print('ERRO: openpyxl não encontrado (rode com python3 -E, não -I).', file=sys.stderr)
        return 1

    problemas_leitura = []
    arvores = ler_export(args.export, problemas_leitura)
    livro = openpyxl.load_workbook(args.planilha, read_only=True, data_only=True)
    try:
        bruto_proj = ler_aba(livro, ABA_PROJECOES, COLUNAS_PROJECOES, problemas_leitura)
        bruto_param = ler_aba(livro, ABA_PARAMETRIZACAO, COLUNAS_PARAMETRIZACAO, problemas_leitura)
    finally:
        livro.close()

    v = Validador(arvores)
    v.problemas.extend(problemas_leitura)
    projecoes, parametrizacao, consolidacao = [], [], []
    if not problemas_leitura:
        projecoes = montar_projecoes(bruto_proj, v)
        parametrizacao = montar_parametrizacao(bruto_param, v)
        consolidacao = montar_consolidacao(v)
        conferir_conjunto(
            (
                (ABA_PROJECOES, projecoes),
                (ABA_PARAMETRIZACAO, parametrizacao),
                (ABA_CONSOLIDACAO, consolidacao),
            ),
            v,
        )

    sql = None
    if not v.problemas:
        try:
            sql = gerar_sql(
                projecoes, parametrizacao, consolidacao, arvores, args.planilha, args.export
            )
        except FormatoInesperado as e:
            v.erro(f'SQL: {e}')

    resumo = (
        f'projecoes={len(projecoes)} parametrizacao={len(parametrizacao)} '
        f'consolidacao={len(consolidacao)} divergencias={len(v.problemas)}'
    )
    if v.problemas:
        for p in v.problemas:
            print(f'ERRO: {p}', file=sys.stderr)
        print(f'Nenhum arquivo gravado. {resumo}', file=sys.stderr)
        return 1

    saida = Path(args.saida)
    saida.parent.mkdir(parents=True, exist_ok=True)
    temporario = saida.with_name(saida.name + '.tmp')
    temporario.write_text(sql, encoding='utf-8')
    temporario.replace(saida)
    print(f'SQL gravado em {saida}')
    print(resumo)
    return 0


if __name__ == '__main__':
    sys.exit(main())
